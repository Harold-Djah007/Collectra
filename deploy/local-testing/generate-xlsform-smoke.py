"""Create a small workbook for testing the XLSForm importer in a demo app."""

import argparse
from pathlib import Path

from openpyxl import Workbook


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', required=True)
    output = Path(parser.parse_args().output).expanduser()
    if output.exists():
        parser.error(f'Refusing to overwrite {output}')
    book = Workbook()
    survey = book.active
    survey.title = 'survey'
    rows = [
        ['type', 'name', 'label', 'required', 'relevant', 'constraint', 'calculation'],
        ['start', 'started', '', '', '', '', ''],
        ['today', 'day', '', '', '', '', ''],
        ['select_one status_list', 'status', 'Equipment status', 'yes', '', '', ''],
        ['text', 'issue', 'Describe the issue', 'yes', "${status} = 'attention'", '', ''],
        ['begin repeat', 'readings', 'Meter readings', '', '', '', ''],
        ['integer', 'reading', 'Reading (zero is allowed)', 'yes', '', '. >= 0', ''],
        ['calculate', 'double_reading', '', '', '', '', '${reading} * 2'],
        ['end repeat', '', '', '', '', '', ''],
        ['calculate', 'total', '', '', '', '', 'sum(${double_reading})'],
        ['end', 'finished', '', '', '', '', ''],
    ]
    for row in rows:
        survey.append(row)
    choices = book.create_sheet('choices')
    for row in [['list_name', 'name', 'label'], ['status_list', 'ok', 'Working normally'],
                ['status_list', 'attention', 'Needs attention'], ['status_list', 'not_done', 'Not checked']]:
        choices.append(row)
    settings = book.create_sheet('settings')
    settings.append(['form_title', 'form_id', 'default_language', 'version'])
    settings.append(['Collectra import demo', 'collectra_import_smoke', 'en', '1'])
    output.parent.mkdir(parents=True, exist_ok=True)
    book.save(output)
    print(f'Wrote {output}. Import it into a separate demonstration application.')


if __name__ == '__main__':
    main()
